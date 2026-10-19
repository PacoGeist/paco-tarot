import re, os
import json, glob
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
D=os.path.join(os.path.dirname(os.path.abspath(__file__)), 'data')+'/'
OUT=os.environ.get('OUT', 'Tarot-Quiz-Content.xlsx')
M=[c for i in range(1,7) for c in json.load(open(f'{D}tarot-meanings-{i}.json'))]
S=[c for i in range(1,7) for c in json.load(open(f'{D}tarot-symbols-{i}.json'))]
CTX=[('past','Looking back'),('present','Here and now'),('future','What is unfolding'),('challenge','The challenge'),('advice','Guidance'),('self','You'),('others','Those around you'),('hopes','Hopes & fears'),('love','Love & the heart'),('work','Work, money & craft'),('spirit','Spirit'),('release','What to release'),('shadow','The hidden side'),('foundation','The foundation'),('gift','The gift')]
F=Font(name='Arial',size=10); B=Font(name='Arial',size=10,bold=True); H=Font(name='Arial',size=10,bold=True,color='FFFFFF')
HF=PatternFill('solid',fgColor='4B2E6F'); LOCK=PatternFill('solid',fgColor='E7E3EE'); NEW=PatternFill('solid',fgColor='FFF2B3')
W=Alignment(wrap_text=True,vertical='top'); line=Side(style='thin',color='CCCCCC'); BD=Border(bottom=line)
wb=Workbook()
def sheet(title,heads,widths,rows,locked):
  ws=wb.create_sheet(title); ws.append(heads)
  for i,w in enumerate(widths): ws.column_dimensions[chr(65+i)].width=w
  for c in ws[1]: c.font=H; c.fill=HF; c.alignment=W
  for r in rows: ws.append(r)
  for row in ws.iter_rows(min_row=2):
    for c in row:
      c.font=F; c.alignment=W; c.border=BD
      if c.column in locked: c.fill=LOCK
  ws.freeze_panes='B2' if 1 in locked else 'A2'; ws.auto_filter.ref=ws.dimensions
  return ws
# --- instructions
ws=wb.active; ws.title='How to edit'; ws.column_dimensions['A'].width=110
txt=[('Tarot School quiz content',True),
('Every quiz question is built from the text in this workbook. Change the text here and the questions, lessons and card pages all change with it.',False),('',False),
('Colours',True),
('Grey cells: card names, positions and IDs. Please don\'t change these; the app uses them to match each row to its card.',False),
('White cells: edit freely.',False),
('Yellow cells (Question Wording sheet): type your new wording here. Leave blank to keep the current wording.',False),('',False),
('The sheets',True),
('Cards: one row per card, with keywords, meaning, picture description, affirmation, element, astrology, numerology and yes/no.',False),
('Positions: how each card reads in each kind of spread position (15 per card).',False),
('Symbols: the features on each card and what they mean. You may add a row (type the card name exactly as in the grey column) or delete a row.',False),
('All School Questions: every question Tarot School can ask, card by card, with the correct answer. This sheet is for reading. To change a question, edit the text on Cards, Positions or Symbols, or its wording on Question Wording.',False),
('Question Wording: the 20 question patterns. {card} is replaced by the card\'s name, e.g. "Which keywords belong to {card}?"',False),('',False),
('Giveaway words: a question never shows the words that name its answer. The app swaps them for vaguer words by itself ("The mountains stand for…" becomes "They stand for…", "nine golden cups" becomes "several golden vessels", "In love, …" becomes "In matters of affection, …"). So you can write the Cards, Positions and Symbols text naturally; the All School Questions sheet shows the result.',False),('',False),
('How quizzes pick questions: each card has 3 ways to ask about every fact (Q01–Q20). A lesson quiz asks about the facts it just taught; the final card quiz picks 15; group quizzes pick 2–5 per card, favouring ones you got wrong or haven\'t seen. Wrong answer choices are drawn at random from other cards.',False),('',False),
('Tips so the questions keep working',True),
('Keywords: separate with commas, e.g. "beginnings, innocence, faith". Keep at least 3 per card.',False),
('Yes/No: must be Yes, No or Maybe.',False),
('Meaning and positions: write full sentences ending in a full stop. Questions quote single sentences, so each sentence should make sense on its own.',False),
('Symbol meanings: the first sentence is what shows up as a quiz answer, so make it a clear summary.',False),('',False),
('When you\'re done',True),
('Save the file (keep the same name) in your Tarot_App folder and tell Claude. Claude checks every change, loads it into the app, tests the quizzes, and includes it in the next update.',False)]
for t,b in txt:
  ws.append([t]); c=ws.cell(ws.max_row,1); c.font=Font(name='Arial',size=14 if ws.max_row==1 else 10,bold=b); c.alignment=Alignment(wrap_text=True)
# --- cards
sheet('Cards',['Card','Keywords (comma separated)','Meaning','Picture description (imagery)','Affirmation','Element','Astrology','Numerology','Yes/No'],
 [24,34,70,70,40,10,22,30,9],
 [[c['name'],', '.join(c['keywords']),c['meaning'],c['imagery'],c['affirmation'],c['element'],c['astrology'],c['numerology'],c['yesNo']] for c in M],{1})
dv=DataValidation(type='list',formula1='"Yes,No,Maybe"',allow_blank=False); wb['Cards'].add_data_validation(dv); dv.add(f'I2:I{len(M)+1}')
for r in range(2,len(M)+2): wb['Cards'].row_dimensions[r].height=150
# --- positions
sheet('Positions',['Card','Position','What it means in this position'],[24,22,100],
 [[c['name'],lab,c['ctx'][k]] for c in M for k,lab in CTX],{1,2})
# --- symbols
sheet('Symbols',['Card','Feature','Where on the card','Meaning'],[24,40,40,80],
 [[c['name'],s['feature'],s['where'],s['meaning']] for c in S for s in c['symbols']],{1})
# --- wording
Q=[('Q01','Recognise the card','Picture of the card; choose its name','Which card is this?'),
('Q02','Recognise the card','Four pictures; choose the right one','Which of these is {card}?'),
('Q03','Recognise the card','Quotes the picture description; choose the card','Which card shows this scene?'),
('Q04','Keywords','Choose the right set of 3 keywords','Which keywords belong to {card}?'),
('Q05','Keywords','Quotes 3 keywords; choose the card','Which card is known for these keywords?'),
('Q06','Keywords','Choose the one real keyword','Which of these is a keyword of {card}?'),
('Q07','Meaning','Choose the sentence from its meaning','What is the heart of {card}’s message?'),
('Q08','Meaning','Quotes a sentence of the meaning; choose the card','Which card carries this message?'),
('Q09','Meaning','Quotes the affirmation; choose the card','Which card’s affirmation is this?'),
('Q10','Meaning','Used instead of Q09 when the affirmation names the card','Which card is this describing?'),
('Q11','Element & astrology','Choose the element','Which element is linked to {card}?'),
('Q12','Element & astrology','Choose the astrology','Which astrological link belongs to {card}?'),
('Q13','Element & astrology','Choose Yes, No or Maybe','Asked a yes-or-no question, how does {card} usually answer?'),
('Q14','Symbols','Quotes a feature; choose its meaning','On {card}, what does this feature mean?'),
('Q15','Symbols','Quotes a meaning; choose the feature','Which feature of {card} carries this meaning?'),
('Q16','Symbols','Shows the card and where to look; choose the feature','Look at {card}. Which feature is found here?'),
('Q17','Symbols','Used instead of Q16 when no "where" is given','Which of these features appears on {card}?'),
('Q18','Positions','Names the position; choose what the card suggests','{card} falls in a “{position}” position. What does it suggest?'),
('Q19','Positions','Quotes a reading; choose the position','This is how {card} reads in one kind of position. Which position?'),
('Q20','Positions','Names the position and quotes a reading; choose the card','Which card, falling in a “{position}” position, would say this?')]
ws=sheet('Question Wording',['ID','Topic','How the question works','Current wording','Your new wording','Notes for Claude'],[7,20,46,52,52,40],[list(q)+['',''] for q in Q],{1,2,3,4})
for r in range(2,len(Q)+2):
  for col in (5,6): ws.cell(r,col).fill=NEW

# --- every school question, made by the app's own code (tools/dump_questions.js)
import subprocess, os
Q_ALL = json.loads(subprocess.run(['node', os.path.join(os.path.dirname(os.path.abspath(__file__)), 'dump_questions.js')], capture_output=True, text=True, check=True).stdout)
KW = {c['name']: c['keywords'] for c in M}
def qid(r):
  k, v, p = r['fact'].split(':')[0], r['v'], r['prompt']
  if k == 'id': return ['Q01', 'Q02', 'Q03'][v]
  if k == 'kw': return ['Q04', 'Q05', 'Q06'][v]
  if k == 'mean': return ['Q07', 'Q08', 'Q09' if 'affirmation' in p else 'Q10'][v]
  if k == 'attr': return 'Q11' if 'element' in p else 'Q12' if 'astrolog' in p else 'Q13'
  if k == 'sym': return ['Q14', 'Q15', 'Q16' if p.startswith('Look at') else 'Q17'][v]
  return ['Q18', 'Q19', 'Q20'][v]
rows = []
for r in Q_ALL:
  q = qid(r); shown = r['quote']; ans = r['answer']
  if q == 'Q01': shown = '(picture of the card, title hidden)'
  if q == 'Q02': shown, ans = '(four card pictures)', 'the picture of ' + ans
  if q in ('Q16', 'Q17'): shown = '(picture) ' + shown if shown else '(picture)'
  if q == 'Q04': ans = 'any 3 of: ' + ans + '  (picked at random each time)'
  if q == 'Q05': shown = shown + '  (3 picked at random each time)'
  rows.append([r['card'], r['step'], q, r['prompt'], shown, ans])
ws=sheet('All School Questions',['Card','Lesson quiz','Question ID','Question','What is shown or quoted','Correct answer'],[24,22,9,48,60,50],rows,{1,2,3})
print('questions',len(rows))
wb.move_sheet('All School Questions',offset=-4)
wb.save(OUT)
print(len(M),sum(len(c['symbols']) for c in S))
